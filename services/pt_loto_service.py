import unicodedata
from datetime import datetime, date

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

from models import db, TreinamentoPTLoto


def _norm(valor):
    """Normaliza texto: sem acentos, maiusculo, sem espacos extras."""
    if valor is None:
        return ""
    texto = unicodedata.normalize("NFKD", str(valor))
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    return " ".join(texto.upper().split())


def _texto(valor):
    """Converte celula em texto limpo (1.0 vira '1')."""
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


def _data(valor):
    """Converte celula em date. Retorna None se invalida."""
    if valor is None or valor == "":
        return None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    if isinstance(valor, (int, float)):
        try:
            return from_excel(valor).date()
        except Exception:
            return None
    texto = str(valor).strip().split(" ")[0]   # remove hora, se houver
    for formato in ("%d/%m/%Y", "%d/%m/%y", "%Y-%m-%d", "%d-%m-%Y",
                    "%d-%m-%y", "%d.%m.%Y", "%d.%m.%y"):
        try:
            return datetime.strptime(texto, formato).date()
        except ValueError:
            continue
    return None


def _achar_cabecalho(ws):
    """Procura a linha de cabecalho nas 15 primeiras linhas. Retorna (linha, colunas)."""
    for idx, row in enumerate(ws.iter_rows(min_row=1, max_row=15, values_only=True), start=1):
        normalizados = [_norm(c) for c in row]
        if any(c.startswith("NOME") for c in normalizados) and \
           any(c.startswith("DATA") for c in normalizados):
            colunas = {}
            for i, c in enumerate(normalizados):
                if "INSTRUTOR" in c:
                    colunas.setdefault("instrutor", i)
                elif c.startswith("NUM"):
                    colunas.setdefault("numero", i)
                elif c.startswith("NOME"):
                    colunas.setdefault("nome", i)
                elif c.startswith("MATRICULA") or c.startswith("SSO"):
                    colunas.setdefault("matricula", i)
                elif c.startswith("EMPRESA"):
                    colunas.setdefault("empresa", i)
                elif c.startswith("DATA") and "VALID" not in c:
                    colunas.setdefault("data", i)
            if "nome" in colunas and "data" in colunas:
                return idx, colunas
    return None, {}


def ler_planilha(stream):
    """
    Le a planilha (procura o cabecalho em todas as abas).
    Retorna (linhas, erros). Cada linha traz o numero da linha no Excel.
    """
    wb = load_workbook(stream, data_only=True)

    ws = None
    linha_cab = None
    colunas = {}
    for planilha in wb.worksheets:
        linha_cab, colunas = _achar_cabecalho(planilha)
        if linha_cab:
            ws = planilha
            break

    if ws is None:
        raise ValueError(
            "Cabecalho nao encontrado em nenhuma aba. A planilha precisa ter "
            "as colunas NOME e DATA DO TREINAMENTO."
        )

    def pegar(row, chave):
        i = colunas.get(chave)
        if i is None or i >= len(row):
            return None
        return row[i]

    linhas = []
    erros = []

    for num_linha, row in enumerate(
        ws.iter_rows(min_row=linha_cab + 1, values_only=True), start=linha_cab + 1
    ):
        if all(c is None or str(c).strip() == "" for c in row):
            continue  # linha vazia

        nome = _texto(pegar(row, "nome"))
        if not nome:
            erros.append(f"Linha {num_linha}: nome vazio.")
            continue

        bruto = pegar(row, "data")
        data_trein = _data(bruto)
        if data_trein is None:
            erros.append(
                f"Linha {num_linha} ({nome}): data do treinamento invalida "
                f"(valor lido: '{_texto(bruto)}')."
            )
            continue

        linhas.append({
            "linha": num_linha,
            "numero": _texto(pegar(row, "numero")),
            "nome": nome,
            "matricula": _texto(pegar(row, "matricula")),
            "empresa": _texto(pegar(row, "empresa")),
            "data_treinamento": data_trein,
            "instrutor": _texto(pegar(row, "instrutor")),
        })

    return linhas, erros


def _chave(nome, matricula, data):
    return (_norm(nome), (matricula or "").strip().upper(), data)


def gravar_linhas(linhas, substituir=False):
    """
    Grava as linhas no banco. Retorna (inseridos, duplicados).
    Duplicado = mesmo nome + mesma matricula/SSO + mesma data de treinamento.
    'duplicados' e uma lista de textos explicando cada linha ignorada.
    """
    if substituir:
        TreinamentoPTLoto.query.delete()
        no_banco = set()
    else:
        no_banco = {
            _chave(r.nome, r.matricula, r.data_treinamento)
            for r in TreinamentoPTLoto.query.all()
        }

    na_planilha = set()
    inseridos = 0
    duplicados = []

    for l in linhas:
        dados = {k: v for k, v in l.items() if k != "linha"}
        chave = _chave(l["nome"], l["matricula"], l["data_treinamento"])

        if chave in no_banco:
            duplicados.append(
                f"Linha {l['linha']} ({l['nome']}, {l['data_treinamento'].strftime('%d/%m/%Y')}): "
                f"ja existe no sistema."
            )
            continue
        if chave in na_planilha:
            duplicados.append(
                f"Linha {l['linha']} ({l['nome']}, {l['data_treinamento'].strftime('%d/%m/%Y')}): "
                f"repetida na planilha."
            )
            continue

        na_planilha.add(chave)
        db.session.add(TreinamentoPTLoto(**dados))
        inseridos += 1

    db.session.commit()
    return inseridos, duplicados
